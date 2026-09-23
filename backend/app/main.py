# from io import BytesIO
# import os
# import subprocess
# import tempfile

# from docx import Document
# from fastapi import FastAPI, File, HTTPException, UploadFile
# from fastapi.middleware.cors import CORSMiddleware
# from fastapi.responses import JSONResponse
# from pypdf import PdfReader

# from .services.analyzer import analyze

# app = FastAPI(title="Nexora Career Intelligence")
# app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"], allow_origin_regex=r"https?://(localhost|127\.0\.0\.1):\d+", allow_credentials=True, allow_methods=["*"], allow_headers=["*"])
# current = {}

# @app.exception_handler(Exception)
# async def unexpected_error(_, exc):
#     print(f"Unexpected server error: {type(exc).__name__}: {exc}")
#     return JSONResponse(status_code=500, content={"detail": "Resume analysis failed on the server. Check the backend terminal."})

# def read_file(content: bytes, filename: str) -> str:
#     extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
#     if extension in {"txt", "md"}:
#         return content.decode("utf-8", errors="ignore")
#     if extension == "pdf":
#         return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
#     if extension == "docx":
#         return "\n".join(paragraph.text for paragraph in Document(BytesIO(content)).paragraphs)
#     if extension == "doc":
#         temp_path = ""
#         try:
#             with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as temp_file:
#                 temp_file.write(content)
#                 temp_path = temp_file.name
#             return subprocess.check_output(["antiword", temp_path], text=True, stderr=subprocess.DEVNULL)
#         except FileNotFoundError as exc:
#             raise HTTPException(415, "Legacy .doc needs antiword. Upload PDF, DOCX, or TXT instead.") from exc
#         except subprocess.CalledProcessError as exc:
#             raise HTTPException(422, "This .doc file could not be read.") from exc
#         finally:
#             if temp_path and os.path.exists(temp_path):
#                 os.unlink(temp_path)
#     raise HTTPException(415, "Supported formats: PDF, DOCX, DOC, TXT, and MD.")

# @app.get("/health")
# def health():
#     return {"status": "ok"}

# @app.post("/api/resume/upload")
# async def upload_resume(file: UploadFile = File(...)):
#     global current
#     filename = file.filename or "resume"
#     content = await file.read()
#     if not content:
#         raise HTTPException(400, "The uploaded file is empty.")
#     text = read_file(content, filename)
#     if len(text.strip()) < 30:
#         raise HTTPException(422, "No readable resume text was found. Try a text-based PDF, DOCX, or TXT file.")
#     current = analyze(text, filename)
#     return current

# @app.get("/api/profile")
# def profile():
#     return current

# @app.get("/api/careers")
# def careers():
#     return current.get("careers", [])

# @app.get("/api/skills/gap")
# def gaps():
#     return {"career": current.get("topCareer", {}).get("name"), "gaps": current.get("gaps", [])}

# @app.get("/api/jobs")
# def jobs():
#     return current.get("jobs", [])

# @app.get("/api/career-twin")
# def career_twin():
#     return current

# @app.post("/api/career/simulate")
# async def simulate(payload: dict):
#     skill = str(payload.get("skill", "")).strip()
#     if not skill:
#         raise HTTPException(400, "Enter a skill to simulate.")
#     simulated = []
#     for career in current.get("careers", []):
#         boost = 14 if skill in career["gaps"] else 5 if skill in career["matchingSkills"] else 0
#         simulated.append({**career, "score": min(99, career["score"] + boost)})
#     return {"addedSkill": skill, "careers": sorted(simulated, key=lambda item: item["score"], reverse=True)}

# @app.post("/api/companion/chat")
# async def chat(payload: dict):
#     question = str(payload.get("message", "")).lower()
#     top_career = current.get("topCareer", {})
#     gaps_list = current.get("gaps", [])
#     skills = [item["name"] for item in current.get("skills", [])][:4]
#     answer = f"Your profile currently points most strongly toward {top_career.get('name', 'a clearer direction')}. That is supported by {', '.join(skills) or 'the evidence visible in your resume'}. A high-impact next move is {gaps_list[0] if gaps_list else 'turning your strongest work into a portfolio case study'}."
#     if "gap" in question:
#         answer += f" Main gaps: {', '.join(gaps_list[:4]) or 'no major gap for the selected path'}."
#     return {"reply": answer}

from io import BytesIO
import os
import subprocess
import tempfile

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pypdf import PdfReader

from .services.analyzer import analyze


app = FastAPI(title="Nexora Career Intelligence")


# =========================================================
# CORS
# =========================================================
# This allows your deployed frontend to communicate
# with the FastAPI backend.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# Temporary in-memory storage
# =========================================================
current = {}


# =========================================================
# Global error handler
# =========================================================
@app.exception_handler(Exception)
async def unexpected_error(_, exc):
    print(
        f"Unexpected server error: "
        f"{type(exc).__name__}: {exc}"
    )

    return JSONResponse(
        status_code=500,
        content={
            "detail": (
                "Resume analysis failed on the server. "
                "Check the backend terminal."
            )
        },
    )


# =========================================================
# Resume file reader
# =========================================================
def read_file(content: bytes, filename: str) -> str:
    extension = (
        filename.rsplit(".", 1)[-1].lower()
        if "." in filename
        else ""
    )

    # TXT / Markdown
    if extension in {"txt", "md"}:
        return content.decode(
            "utf-8",
            errors="ignore"
        )

    # PDF
    if extension == "pdf":
        reader = PdfReader(BytesIO(content))

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    # DOCX
    if extension == "docx":
        document = Document(BytesIO(content))

        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

    # Old DOC format
    if extension == "doc":
        temp_path = ""

        try:
            with tempfile.NamedTemporaryFile(
                suffix=".doc",
                delete=False
            ) as temp_file:

                temp_file.write(content)
                temp_path = temp_file.name

            return subprocess.check_output(
                ["antiword", temp_path],
                text=True,
                stderr=subprocess.DEVNULL,
            )

        except FileNotFoundError as exc:
            raise HTTPException(
                status_code=415,
                detail=(
                    "Legacy .doc needs antiword. "
                    "Upload PDF, DOCX, or TXT instead."
                ),
            ) from exc

        except subprocess.CalledProcessError as exc:
            raise HTTPException(
                status_code=422,
                detail="This .doc file could not be read.",
            ) from exc

        finally:
            if temp_path and os.path.exists(temp_path):
                os.unlink(temp_path)

    raise HTTPException(
        status_code=415,
        detail=(
            "Supported formats: "
            "PDF, DOCX, DOC, TXT, and MD."
        ),
    )


# =========================================================
# Health check
# =========================================================
@app.get("/health")
def health():
    return {
        "status": "ok"
    }


# =========================================================
# Resume Upload + Analysis
# =========================================================
@app.post("/api/resume/upload")
async def upload_resume(
    file: UploadFile = File(...)
):
    global current

    filename = file.filename or "resume"

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    text = read_file(
        content,
        filename
    )

    if len(text.strip()) < 30:
        raise HTTPException(
            status_code=422,
            detail=(
                "No readable resume text was found. "
                "Try a text-based PDF, DOCX, or TXT file."
            ),
        )

    current = analyze(
        text,
        filename
    )

    return current


# =========================================================
# Profile
# =========================================================
@app.get("/api/profile")
def profile():
    return current


# =========================================================
# Careers
# =========================================================
@app.get("/api/careers")
def careers():
    return current.get(
        "careers",
        []
    )


# =========================================================
# Skill Gap
# =========================================================
@app.get("/api/skills/gap")
def gaps():
    return {
        "career": current.get(
            "topCareer",
            {}
        ).get(
            "name"
        ),
        "gaps": current.get(
            "gaps",
            []
        ),
    }


# =========================================================
# Jobs
# =========================================================
@app.get("/api/jobs")
def jobs():
    return current.get(
        "jobs",
        []
    )


# =========================================================
# Career Twin
# =========================================================
@app.get("/api/career-twin")
def career_twin():
    return current


# =========================================================
# Career Simulation
# =========================================================
@app.post("/api/career/simulate")
async def simulate(payload: dict):

    skill = str(
        payload.get(
            "skill",
            ""
        )
    ).strip()

    if not skill:
        raise HTTPException(
            status_code=400,
            detail="Enter a skill to simulate.",
        )

    simulated = []

    for career in current.get(
        "careers",
        []
    ):

        boost = (
            14
            if skill in career.get("gaps", [])
            else 5
            if skill in career.get("matchingSkills", [])
            else 0
        )

        simulated.append(
            {
                **career,
                "score": min(
                    99,
                    career["score"] + boost
                ),
            }
        )

    return {
        "addedSkill": skill,
        "careers": sorted(
            simulated,
            key=lambda item: item["score"],
            reverse=True,
        ),
    }


# =========================================================
# AI Career Companion Chat
# =========================================================
@app.post("/api/companion/chat")
async def chat(payload: dict):

    question = str(
        payload.get(
            "message",
            ""
        )
    ).lower()

    top_career = current.get(
        "topCareer",
        {}
    )

    gaps_list = current.get(
        "gaps",
        []
    )

    skills = [
        item["name"]
        for item in current.get(
            "skills",
            []
        )[:4]
    ]

    answer = (
        f"Your profile currently points most strongly "
        f"toward {top_career.get('name', 'a clearer direction')}. "
        f"That is supported by "
        f"{', '.join(skills) or 'the evidence visible in your resume'}. "
        f"A high-impact next move is "
        f"{gaps_list[0] if gaps_list else 'turning your strongest work into a portfolio case study'}."
    )

    if "gap" in question:
        answer += (
            f" Main gaps: "
            f"{', '.join(gaps_list[:4]) or 'no major gap for the selected path'}."
        )

    return {
        "reply": answer
    }